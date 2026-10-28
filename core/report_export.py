"""ส่งออกรายงานอัตราการใช้งานเป็นไฟล์ Excel (.xlsx) และ CSV (TOR 4.5)

รับข้อมูลรายชั่วโมงที่แปลงเป็นเวลาท้องถิ่นแล้ว — ไม่ยุ่งกับฐานข้อมูล ทดสอบได้ตรงๆ
"""
from __future__ import annotations

import csv
import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime

THAI_MONTHS = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
THAI_DAYS = ["จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์", "เสาร์", "อาทิตย์"]


def safe_text(value):
    """กัน formula injection: ข้อความที่ขึ้นต้นด้วย = + - @ จะถูก Excel ตีความเป็นสูตร (เช่น ชื่อโซนที่ผู้ใช้ตั้ง)
    ใส่ ' นำหน้าให้เป็นข้อความธรรมดา — ตัวเลขไม่แตะ"""
    if isinstance(value, str) and value != "-" and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def put(ws, row: int, column: int, value):
    """เขียนเซลล์ Excel โดยไม่ให้ข้อความที่ขึ้นต้นด้วย = กลายเป็นสูตร (openpyxl ถือว่า "=..." คือสูตร)"""
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def thai_date(d: date) -> str:
    """28 ต.ค. 2569 (ปี พ.ศ. ตามแบบเอกสารราชการ)"""
    return f"{d.day} {THAI_MONTHS[d.month - 1]} {d.year + 543}"


def thai_datetime(dt: datetime) -> str:
    return f"{thai_date(dt.date())} {dt:%H:%M} น."


@dataclass
class ZoneReport:
    code: str
    name_th: str
    capacity: int  # จำนวนช่องที่นับยอดได้ตอนออกรายงาน
    # (เวลาเริ่มชั่วโมงแบบท้องถิ่น, เฉลี่ย %, สูงสุด %, จำนวนตัวอย่างรายนาที)
    hours: list[tuple[datetime, float, float, int]] = field(default_factory=list)


@dataclass
class ReportMeta:
    site_name: str
    date_from: date
    date_to: date
    generated_at: datetime  # เวลาท้องถิ่น
    generated_by: str


def daily_rows(zone: ZoneReport) -> list[tuple[date, float, float, int]]:
    """(วันที่, เฉลี่ย %, สูงสุด %, จำนวนชั่วโมงที่มีข้อมูล) เรียงตามวันที่"""
    by_day: dict[date, list[tuple[float, float]]] = defaultdict(list)
    for ts, avg, mx, _n in zone.hours:
        by_day[ts.date()].append((avg, mx))
    return [
        (d, sum(v[0] for v in vals) / len(vals), max(v[1] for v in vals), len(vals))
        for d, vals in sorted(by_day.items())
    ]


def heatmap(zone: ZoneReport) -> dict[tuple[int, int], float]:
    """เฉลี่ย % ตาม (วันในสัปดาห์ 0=จันทร์, ชั่วโมง)"""
    buckets: dict[tuple[int, int], list[float]] = defaultdict(list)
    for ts, avg, _mx, _n in zone.hours:
        buckets[(ts.weekday(), ts.hour)].append(avg)
    return {key: sum(vals) / len(vals) for key, vals in buckets.items()}


def zone_summary(zone: ZoneReport) -> dict:
    """ตัวเลขสรุปของโซนในช่วงรายงาน (ไม่มีข้อมูล = None)"""
    if not zone.hours:
        return {"hours": 0, "avg": None, "max": None, "busiest_day": None, "busiest_hour": None}
    avg = sum(h[1] for h in zone.hours) / len(zone.hours)
    mx = max(h[2] for h in zone.hours)
    days = daily_rows(zone)
    busiest_day = max(days, key=lambda r: r[1])[0]
    by_hour: dict[int, list[float]] = defaultdict(list)
    for ts, a, _m, _n in zone.hours:
        by_hour[ts.hour].append(a)
    busiest_hour = max(by_hour.items(), key=lambda kv: sum(kv[1]) / len(kv[1]))[0]
    return {"hours": len(zone.hours), "avg": avg, "max": mx, "busiest_day": busiest_day, "busiest_hour": busiest_hour}


def hour_label(hour: int) -> str:
    return f"{hour:02d}:00–{hour:02d}:59"


def build_csv(zones: list[ZoneReport]) -> bytes:
    """CSV รายชั่วโมง (ละเอียดที่สุด) มี BOM ให้ Excel เปิดภาษาไทยได้ถูกต้อง"""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["วันที่", "เวลา", "รหัสโซน", "ชื่อโซน", "อัตราใช้งานเฉลี่ย (%)", "อัตราใช้งานสูงสุด (%)", "จำนวนตัวอย่าง (นาที)"])
    rows = sorted(((ts, z, avg, mx, n) for z in zones for ts, avg, mx, n in z.hours), key=lambda r: (r[0], r[1].code))
    for ts, z, avg, mx, n in rows:
        writer.writerow([ts.strftime("%Y-%m-%d"), hour_label(ts.hour), safe_text(z.code), safe_text(z.name_th), f"{avg:.1f}", f"{mx:.1f}", n])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _heat_fill(pct: float) -> str:
    """สีพื้นหลังเซลล์ heatmap: เขียวอ่อน (ว่าง) -> เหลือง -> แดง (หนาแน่น)"""
    t = max(0.0, min(100.0, pct)) / 100
    if t < 0.5:
        a, b, k = (0xDC, 0xFC, 0xE7), (0xFE, 0xF0, 0x8A), t / 0.5
    else:
        a, b, k = (0xFE, 0xF0, 0x8A), (0xF8, 0x71, 0x71), (t - 0.5) / 0.5
    return "".join(f"{round(a[i] + (b[i] - a[i]) * k):02X}" for i in range(3))


def build_xlsx(meta: ReportMeta, zones: list[ZoneReport]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8EEFB")
    thin = Side(style="thin", color="C9D2E3")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    def header(ws, row: int, labels: list[str]) -> None:
        for col, label in enumerate(labels, start=1):
            cell = ws.cell(row=row, column=col, value=label)
            cell.font = bold
            cell.fill = head_fill
            cell.border = box
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def widths(ws, values: list[int]) -> None:
        for i, w in enumerate(values, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ---------- สรุป ----------
    ws = wb.active
    ws.title = "สรุป"
    ws["A1"] = "รายงานอัตราการใช้งานลานจอดรถ"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = meta.site_name
    ws["A2"].font = Font(size=12)
    ws["A3"] = f"ช่วงข้อมูล {thai_date(meta.date_from)} ถึง {thai_date(meta.date_to)}"
    ws["A4"] = f"ออกรายงานเมื่อ {thai_datetime(meta.generated_at)} โดย {meta.generated_by}"
    ws["A4"].font = Font(color="6B7486")
    header(ws, 6, ["รหัสโซน", "ชื่อโซน", "จำนวนช่องจอด", "ชั่วโมงที่มีข้อมูล", "อัตราใช้งานเฉลี่ย (%)",
                   "อัตราใช้งานสูงสุด (%)", "วันที่ใช้งานมากที่สุด", "ช่วงเวลาที่หนาแน่นที่สุด"])
    row = 7
    for z in zones:
        s = zone_summary(z)
        values = [
            z.code, z.name_th, z.capacity, s["hours"],
            round(s["avg"], 1) if s["avg"] is not None else "ไม่มีข้อมูล",
            round(s["max"], 1) if s["max"] is not None else "ไม่มีข้อมูล",
            thai_date(s["busiest_day"]) if s["busiest_day"] else "-",
            hour_label(s["busiest_hour"]) + " น." if s["busiest_hour"] is not None else "-",
        ]
        for col, v in enumerate(values, start=1):
            cell = put(ws, row, col, v)
            cell.border = box
            if col in (5, 6) and isinstance(v, float):
                cell.number_format = "0.0"
        row += 1
    notes = [
        "หมายเหตุ",
        "• อัตราใช้งาน = จำนวนช่องที่มีรถจอด ÷ จำนวนช่องที่นับยอดได้ (ไม่รวมช่องสงวน และช่องที่ไม่ทราบสถานะ) คำนวณทุก 1 นาที แล้วเฉลี่ยรายชั่วโมง",
        "• เวลาทั้งหมดเป็นเวลาประเทศไทย ชั่วโมงที่ไม่มีข้อมูล (เช่น กล้องขาดการเชื่อมต่อ) จะไม่ถูกนับในค่าเฉลี่ย",
        "• ระบบเก็บสถิติรายนาที 30 วัน และรายชั่วโมง 2 ปี ตามข้อกำหนด",
    ]
    row += 1
    for i, text in enumerate(notes):
        cell = ws.cell(row=row + i, column=1, value=text)
        cell.font = bold if i == 0 else Font(color="4B5565")
    widths(ws, [10, 28, 13, 14, 16, 16, 20, 22])
    ws.row_dimensions[6].height = 32

    # ---------- รายวัน ----------
    ws = wb.create_sheet("รายวัน")
    header(ws, 1, ["วันที่", "วัน", "รหัสโซน", "ชื่อโซน", "อัตราใช้งานเฉลี่ย (%)", "อัตราใช้งานสูงสุด (%)", "ชั่วโมงที่มีข้อมูล"])
    r = 2
    day_rows = sorted(((d, z, a, m, n) for z in zones for d, a, m, n in daily_rows(z)), key=lambda x: (x[0], x[1].code))
    for d, z, a, m, n in day_rows:
        for col, v in enumerate([d, THAI_DAYS[d.weekday()], z.code, z.name_th, round(a, 1), round(m, 1), n], start=1):
            cell = put(ws, r, col, v)
            if col == 1:
                cell.number_format = "dd/mm/yyyy"
            if col in (5, 6):
                cell.number_format = "0.0"
        r += 1
    widths(ws, [12, 11, 10, 28, 16, 16, 14])
    ws.freeze_panes = "A2"
    if r > 2:
        ws.auto_filter.ref = f"A1:G{r - 1}"

    # ---------- รายชั่วโมง ----------
    ws = wb.create_sheet("รายชั่วโมง")
    header(ws, 1, ["วันที่", "เวลา", "รหัสโซน", "ชื่อโซน", "อัตราใช้งานเฉลี่ย (%)", "อัตราใช้งานสูงสุด (%)", "จำนวนตัวอย่าง (นาที)"])
    r = 2
    hour_rows = sorted(((ts, z, a, m, n) for z in zones for ts, a, m, n in z.hours), key=lambda x: (x[0], x[1].code))
    for ts, z, a, m, n in hour_rows:
        for col, v in enumerate([ts.date(), hour_label(ts.hour), z.code, z.name_th, round(a, 1), round(m, 1), n], start=1):
            cell = put(ws, r, col, v)
            if col == 1:
                cell.number_format = "dd/mm/yyyy"
            if col in (5, 6):
                cell.number_format = "0.0"
        r += 1
    widths(ws, [12, 14, 10, 28, 16, 16, 16])
    ws.freeze_panes = "A2"
    if r > 2:
        ws.auto_filter.ref = f"A1:G{r - 1}"

    # ---------- Heatmap ----------
    ws = wb.create_sheet("Heatmap")
    ws["A1"] = "แผนภาพความหนาแน่น: อัตราใช้งานเฉลี่ย (%) ตามวันในสัปดาห์ และช่วงเวลา"
    ws["A1"].font = Font(bold=True, size=13)
    r = 3
    for z in zones:
        cells = heatmap(z)
        put(ws, r, 1, f"โซน {z.code} · {z.name_th}").font = bold
        r += 1
        header(ws, r, ["วัน \\ เวลา"] + [f"{h:02d}" for h in range(24)])
        for dow in range(7):
            ws.cell(row=r + 1 + dow, column=1, value=THAI_DAYS[dow]).font = bold
            for h in range(24):
                cell = ws.cell(row=r + 1 + dow, column=2 + h)
                cell.border = box
                cell.alignment = center
                if (dow, h) in cells:
                    v = round(cells[(dow, h)])
                    cell.value = v
                    cell.fill = PatternFill("solid", fgColor=_heat_fill(v))
        r += 10
    ws.column_dimensions["A"].width = 12
    for h in range(24):
        ws.column_dimensions[get_column_letter(2 + h)].width = 4.6
    ws.cell(row=r, column=1, value="เซลล์ว่าง = ไม่มีข้อมูลในช่วงเวลานั้น · สีเขียว = ว่าง, สีแดง = หนาแน่น").font = Font(color="6B7486")

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = sheet.title != "สรุป"
        # พิมพ์ได้สวยทันที: แนวนอนสำหรับสรุป/Heatmap, ย่อให้พอดีความกว้างกระดาษ A4
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = "landscape" if sheet.title in ("สรุป", "Heatmap") else "portrait"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        if sheet.title in ("รายวัน", "รายชั่วโมง"):
            sheet.print_title_rows = "1:1"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
