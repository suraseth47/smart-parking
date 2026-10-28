"""ส่งออกรายงาน Excel/CSV และการแปลงเวลา UTC -> เวลาไทยในรายงาน (TOR 4.5)"""
import io
from datetime import date, datetime

from openpyxl import load_workbook

from core.report_export import ReportMeta, ZoneReport, build_csv, build_xlsx, daily_rows, safe_text, zone_summary
from core.reports import build_daily_report, build_heatmap_report, build_hourly_report

META = ReportMeta("ลานจอดทดสอบ", date(2026, 10, 1), date(2026, 10, 2), datetime(2026, 10, 3, 9, 0), "admin1")


def zone(name="ลานหน้าอาคาร"):
    z = ZoneReport(code="A", name_th=name, capacity=8)
    z.hours = [
        (datetime(2026, 10, 1, 9), 80.0, 100.0, 60),
        (datetime(2026, 10, 1, 10), 60.0, 75.0, 60),
        (datetime(2026, 10, 2, 9), 40.0, 50.0, 30),
    ]
    return z


def test_summary_and_daily():
    z = zone()
    s = zone_summary(z)
    assert s["hours"] == 3 and round(s["avg"], 1) == 60.0 and s["max"] == 100.0
    assert s["busiest_day"] == date(2026, 10, 1) and s["busiest_hour"] == 9
    assert daily_rows(z)[0] == (date(2026, 10, 1), 70.0, 100.0, 2)


def test_xlsx_has_all_sheets_and_values():
    wb = load_workbook(io.BytesIO(build_xlsx(META, [zone()])))
    assert wb.sheetnames == ["สรุป", "รายวัน", "รายชั่วโมง", "Heatmap"]
    assert wb["สรุป"]["A7"].value == "A" and wb["สรุป"]["E7"].value == 60.0
    assert wb["รายชั่วโมง"].max_row == 4  # หัวตาราง + 3 ชั่วโมง


def test_xlsx_never_writes_formulas_from_user_text():
    wb = load_workbook(io.BytesIO(build_xlsx(META, [zone(name="=HYPERLINK(\"http://x\")")])))
    cell = wb["สรุป"]["B7"]
    assert cell.data_type == "s" and cell.value.startswith("=")


def test_csv_has_bom_and_guards_formula():
    text = build_csv([zone(name="=cmd")]).decode("utf-8")
    assert text.startswith("﻿") and "'=cmd" in text
    assert safe_text("-") == "-" and safe_text("+1") == "'+1" and safe_text(5) == 5


def test_reports_use_local_time():
    # 02:00 UTC = 09:00 เวลาไทย, 20:00 UTC วันที่ 1 = 03:00 วันที่ 2 เวลาไทย
    rows = [(datetime(2026, 10, 1, 2), 50.0, 60.0, 60), (datetime(2026, 10, 1, 20), 10.0, 20.0, 60)]
    hourly = build_hourly_report(rows, 7)
    assert hourly[9].samples == 60 and hourly[2].samples == 0
    daily = build_daily_report([(ts, a, m) for ts, a, m, _ in rows], 7)
    assert [d.date for d in daily] == ["2026-10-01", "2026-10-02"]
    cells = build_heatmap_report([(ts, a) for ts, a, _, _ in rows], 7)
    assert {(c.day_of_week, c.hour) for c in cells} == {(3, 9), (4, 3)}  # พฤหัส 09:00, ศุกร์ 03:00
