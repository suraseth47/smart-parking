"""ส่งออกตารางจากหน้าเว็บเป็น Excel (ปุ่ม Excel ในเมนูต่างๆ)"""
import io
from datetime import datetime

from openpyxl import load_workbook

from web.routers.export import TableExportIn, build_table_xlsx, safe_filename


def test_safe_filename():
    assert safe_filename("users") == "users"
    assert safe_filename("ผู้ใช้ 2/3") == "2-3"
    assert safe_filename("ทั้งหมด") == "export"


def test_build_table_xlsx():
    data = TableExportIn(
        title="รายชื่อผู้ใช้", filename="users", columns=["ชื่อผู้ใช้", "จำนวน"],
        rows=[["admin1", 3], ["=cmd|calc", None]], meta=["สถานที่: ทดสอบ"],
    )
    wb = load_workbook(io.BytesIO(build_table_xlsx(data, datetime(2026, 10, 28, 15, 0), "admin1")))
    ws = wb.active
    assert ws.title == "รายชื่อผู้ใช้"
    assert ws["A1"].value == "รายชื่อผู้ใช้"
    assert ws["A3"].value == "สถานที่: ทดสอบ"
    assert [c.value for c in ws[5]] == ["ชื่อผู้ใช้", "จำนวน"]  # หัวตารางอยู่แถว 5 (ชื่อ/เวลา/meta/เว้น 1 แถว)
    assert ws["B6"].value == 3
    assert ws["A7"].data_type == "s"  # ข้อความขึ้นต้นด้วย = ต้องไม่กลายเป็นสูตร
