"""Audit log viewer (FR-43) + ส่งออกไฟล์ (TOR 4.5) — Admin เท่านั้น"""
from __future__ import annotations

import csv
import io
import json
from datetime import date as date_type
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from core.audit_labels import ACTION_LABELS, action_label
from core.config import local_utc_offset_hours
from core.db import get_db
from core.models import AuditLog, User
from core.report_export import put, safe_text, thai_datetime
from web.deps import CurrentUser, require_admin

router = APIRouter(prefix="/api/admin/audit", tags=["audit"])

MAX_EXPORT_ROWS = 50000


class AuditLogOut(BaseModel):
    id: int
    username: str | None
    action: str
    action_label: str
    target: str | None
    detail: dict | None
    ts: datetime


class AuditPageOut(BaseModel):
    total: int
    items: list[AuditLogOut]


def _filtered_query(
    action: str | None, username: str | None, from_: date_type | None, to: date_type | None
):
    query = select(AuditLog, User.username).outerjoin(User, User.id == AuditLog.user_id)
    if action:
        query = query.where(AuditLog.action == action)
    if username:
        query = query.where(User.username == username.strip())
    shift = timedelta(hours=local_utc_offset_hours())
    if from_:
        query = query.where(AuditLog.ts >= datetime(from_.year, from_.month, from_.day) - shift)
    if to:
        query = query.where(AuditLog.ts < datetime(to.year, to.month, to.day) + timedelta(days=1) - shift)
    return query


def _row(log: AuditLog, username: str | None) -> dict:
    return {
        "id": log.id,
        "username": username,
        "action": log.action,
        "action_label": action_label(log.action),
        "target": log.target,
        "detail": log.detail,
        "ts": log.ts,
    }


@router.get("/actions")
def list_actions(_admin: CurrentUser = Depends(require_admin)) -> list[dict]:
    """รายการประเภทการกระทำ (ใช้ทำตัวกรองในหน้าเว็บ)"""
    return [{"action": k, "label": v} for k, v in sorted(ACTION_LABELS.items(), key=lambda kv: kv[1])]


@router.get("", response_model=AuditPageOut)
def list_audit_log(
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
    action: str | None = None,
    username: str | None = None,
    from_: date_type | None = Query(None, alias="from"),
    to: date_type | None = None,
    db: Session = Depends(get_db),
    _admin: CurrentUser = Depends(require_admin),
):
    query = _filtered_query(action, username, from_, to)
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rows = db.execute(query.order_by(AuditLog.ts.desc(), AuditLog.id.desc()).offset(offset).limit(limit)).all()
    return {"total": total, "items": [_row(log, username) for log, username in rows]}


@router.get("/export")
def export_audit_log(
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    action: str | None = None,
    username: str | None = None,
    from_: date_type | None = Query(None, alias="from"),
    to: date_type | None = None,
    db: Session = Depends(get_db),
    admin: CurrentUser = Depends(require_admin),
):
    query = _filtered_query(action, username, from_, to).order_by(AuditLog.ts.desc(), AuditLog.id.desc())
    rows = db.execute(query.limit(MAX_EXPORT_ROWS + 1)).all()
    if len(rows) > MAX_EXPORT_ROWS:
        raise HTTPException(422, f"ข้อมูลเกิน {MAX_EXPORT_ROWS:,} รายการ กรุณาเลือกช่วงวันที่ให้แคบลง")

    shift = timedelta(hours=local_utc_offset_hours())
    header = ["วันเวลา", "ผู้ใช้", "การกระทำ", "รหัสการกระทำ", "เป้าหมาย", "รายละเอียด"]
    table = [
        [
            log.ts + shift,
            name or "-",
            action_label(log.action),
            log.action,
            log.target or "",
            json.dumps(log.detail, ensure_ascii=False) if log.detail else "",
        ]
        for log, name in rows
    ]
    stamp = datetime.utcnow() + shift
    filename = f"audit-log-{stamp:%Y%m%d-%H%M}.{format}"

    if format == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(header)
        for r in table:
            writer.writerow([r[0].strftime("%Y-%m-%d %H:%M:%S"), *[safe_text(v) for v in r[1:]]])
        content = ("﻿" + buf.getvalue()).encode("utf-8")
        media = "text/csv; charset=utf-8"
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "ประวัติการใช้งาน"
        ws.append(["ประวัติการใช้งานระบบ (Audit log)"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"ส่งออกเมื่อ {thai_datetime(stamp)} โดย {admin.username} · {len(table):,} รายการ · เวลาเป็นเวลาประเทศไทย"])
        ws.append([])
        ws.append(header)
        for cell in ws[4]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EEFB")
        for i, r in enumerate(table, start=5):
            for col, v in enumerate(r, start=1):
                put(ws, i, col, v)
            ws.cell(row=i, column=1).number_format = "dd/mm/yyyy hh:mm:ss"
        for col, width in zip("ABCDEF", (20, 16, 30, 22, 26, 60)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A5"
        if table:
            ws.auto_filter.ref = f"A4:F{ws.max_row}"
        out = io.BytesIO()
        wb.save(out)
        content = out.getvalue()
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    db.add(AuditLog(user_id=admin.id, action="export_audit", target=filename[:100],
                    detail={"rows": len(table), "format": format, "from": str(from_) if from_ else None, "to": str(to) if to else None}))
    db.commit()
    return Response(
        content,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"},
    )
