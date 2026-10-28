"""ส่งออกตารางที่แสดงบนหน้าเว็บเป็นไฟล์ Excel (ปุ่ม "Excel" ในเมนูต่างๆ)

หน้าเว็บส่ง "ตารางที่ผู้ใช้เห็นอยู่" (หลังกรองแล้ว) มาให้จัดรูปแบบเป็น .xlsx
ข้อมูลมาจาก API ที่ตรวจสิทธิ์ไปแล้ว endpoint นี้แค่จัดรูปแบบไฟล์ และบันทึกว่าใครส่งออกอะไร
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from core.config import local_utc_offset_hours
from core.db import get_db
from core.models import AuditLog
from core.report_export import put, thai_datetime
from web.deps import CurrentUser, get_current_user, verify_csrf

router = APIRouter(prefix="/api/admin/export", tags=["export"])

MAX_ROWS = 50000
MAX_COLS = 40
MAX_CELL = 2000
Cell = str | int | float | bool | None


class TableExportIn(BaseModel):
    title: str = Field(min_length=1, max_length=120)
    filename: str = Field(min_length=1, max_length=80)
    columns: list[str] = Field(min_length=1, max_length=MAX_COLS)
    rows: list[list[Cell]] = Field(max_length=MAX_ROWS)
    meta: list[str] = Field(default_factory=list, max_length=10)  # บรรทัดอธิบาย เช่น ตัวกรองที่ใช้


def safe_filename(name: str) -> str:
    """ชื่อไฟล์ภาษาอังกฤษล้วน (header ของ HTTP ใส่ภาษาไทยตรงๆ ไม่ได้)"""
    cleaned = re.sub(r"[^A-Za-z0-9_-]+", "-", name).strip("-")
    return (cleaned or "export")[:60]


def build_table_xlsx(data: TableExportIn, generated_at: datetime, username: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]\*\?/\\:]", " ", data.title)[:31] or "Sheet1"
    put(ws, 1, 1, data.title).font = Font(bold=True, size=14)
    put(ws, 2, 1, f"ส่งออกเมื่อ {thai_datetime(generated_at)} โดย {username} · {len(data.rows):,} รายการ").font = Font(color="6B7486")
    row = 3
    for line in data.meta:
        put(ws, row, 1, line[:MAX_CELL]).font = Font(color="4B5565")
        row += 1
    head_row = row + 1

    thin = Side(style="thin", color="C9D2E3")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, label in enumerate(data.columns, start=1):
        cell = put(ws, head_row, col, label[:200])
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEFB")
        cell.border = box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [min(max(len(str(c)) + 2, 8), 50) for c in data.columns]
    for r, values in enumerate(data.rows, start=head_row + 1):
        for c, value in enumerate(values[: len(data.columns)], start=1):
            if isinstance(value, str):
                value = value[:MAX_CELL]
            cell = put(ws, r, c, value)
            cell.border = box
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            widths[c - 1] = min(max(widths[c - 1], len(str(value or "")) + 2), 50)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    last_row = head_row + len(data.rows)
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    if data.rows:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(data.columns))}{last_row}"
    ws.print_title_rows = f"{head_row}:{head_row}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if len(data.columns) > 5 else "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@router.post("/table")
def export_table(
    data: TableExportIn,
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
    _csrf: None = Depends(verify_csrf),
):
    if any(len(r) > MAX_COLS for r in data.rows):
        raise HTTPException(422, "จำนวนคอลัมน์เกินกำหนด")
    now_local = datetime.utcnow() + timedelta(hours=local_utc_offset_hours())
    content = build_table_xlsx(data, now_local, user.username)
    filename = f"{safe_filename(data.filename)}-{now_local:%Y%m%d-%H%M}.xlsx"
    db.add(AuditLog(user_id=user.id, action="export_table", target=filename[:100],
                    detail={"title": data.title, "rows": len(data.rows), "format": "xlsx"}))
    db.commit()
    return Response(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"},
    )


class PrintLogIn(BaseModel):
    title: str = Field(min_length=1, max_length=120)
    rows: int = Field(ge=0)
    kind: str = Field(pattern="^(pdf|print)$")


@router.post("/print-log")
def log_print(
    data: PrintLogIn,
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
    _csrf: None = Depends(verify_csrf),
) -> dict:
    """บันทึกการพิมพ์/บันทึก PDF ลงประวัติการใช้งาน (การพิมพ์ทำในเบราว์เซอร์ จึงแจ้งมาที่นี่)"""
    db.add(AuditLog(user_id=user.id, action="export_table", target=data.title[:100],
                    detail={"title": data.title, "rows": data.rows, "format": data.kind}))
    db.commit()
    return {"ok": True}
